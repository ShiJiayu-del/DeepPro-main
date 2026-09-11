#!/usr/bin/env python3
"""Validate and summarize the seven corrected BC-TPro best-val reruns.

The checkpoint inside each run is selected only by internal-validation pixel
IoU.  Models are compared afterwards using the joint Pd/Fa/AUC Pareto rule;
there is no AUC-first rule, weighted score, or unique scalar ranking.
"""

import argparse
import csv
import json
import math
import os
import re
import sys
from copy import copy
from pathlib import Path


REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from runtime_utils import load_checkpoint  # noqa: E402
from tools import analyze_bc_tpro_nongate as joint  # noqa: E402
from tools import run_bc_tpro_bestval as protocol  # noqa: E402


DEFAULT_EXPERIMENT = REPO_ROOT / 'experiments' / protocol.EXPERIMENT_NAME
DEFAULT_LOG_ROOT = REPO_ROOT / 'log' / 'sem_seg'
WORKBOOK_NAME = 'BESTVAL_EXPERIMENT_RESULTS_2026-09-11.xlsx'
LABEL_BY_RUN_ID = {
    run_id: label
    for run_id, _wave, _variant, _gpu, label in protocol.REGISTERED_VARIANTS
}
PROTOCOL_NOTES = (
    '状态要求：7 个作业均有成功 done marker，且 launcher 产物核验通过。',
    '数据集：NUDT-MIRSDT-Noise8.0_FJY；固定 train64/internal-val16；仅 seed47。',
    '每个 epoch 完整验证一次，共 32 次；每个模型按官方逐窗口累计的 internal-val pixel IoU 最大值保存 best_model.pth。',
    '最终 Pd@0.5、Fa@0.5 和 AUC27 必须由该模型的 best_model.pth 在同一 internal-val16 上独立评测。',
    'Pd、Fa、AUC 均从逐序列整数计数重新计算，不信任仅有的汇总小数。',
    '跨模型联合目标为 Pd@0.5 越高、Fa@0.5 越低、AUC27 越高；使用三目标 Pareto 非支配关系。',
    '不存在 AUC 优先规则、加权综合分或唯一标量排名；Pareto 非支配不等于唯一最佳。',
    'best_val_iou 只用于同一训练运行内选择 checkpoint，不参与跨模型三目标 Pareto 比较。',
    '单 seed 不能支持训练随机性方差、统计显著性或稳定性结论。',
    '本表是 internal-val16 结果，不可冒充 official test20 结果。',
)
FLOAT_PATTERN = r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?'


def _configure_launcher(experiment, log_root, launcher=None):
    """Point the protocol validator at explicitly requested read-only roots."""
    experiment = Path(experiment).resolve()
    log_root = Path(log_root).resolve()
    if log_root.name != 'sem_seg':
        raise ValueError('log_root must be the sem_seg directory.')
    runner = launcher or protocol.Launcher()
    runner.experiment = experiment
    runner.save = log_root.parent
    runner.queue = log_root / '_queues' / experiment.name
    runner.status = runner.queue / 'status'
    return runner


def _validation_names(runner):
    path = runner.split('val_sequences.txt')
    if not path.is_file():
        raise FileNotFoundError('Missing internal-validation split: %s' % path)
    raw = path.read_text(encoding='utf-8').splitlines()
    if any(not line.strip() for line in raw):
        raise ValueError('Internal-validation split contains blank lines: %s' % path)
    names = [line.strip() for line in raw]
    expected = list(joint.legacy.VAL_NAMES)
    if names != expected:
        raise ValueError('Internal-validation split is not the frozen ordered val16 set.')
    return names


def _validated_curve(payload, val_names, context):
    curve = joint.legacy.common.validate_curve_counts(payload, val_names, context)
    joint.legacy.common.validate_payload_metrics(
        payload,
        curve,
        context,
        expected_inference_amp=False,
        validate_pixel_diagnostics=False,
    )
    return curve


def _logged_validation_selection(training_log):
    """Recover and internally cross-check the 32 logged validation decisions."""
    text = Path(training_log).read_text(encoding='utf-8')
    markers = list(re.finditer(r'---- EPOCH ([0-9]{3}) EVALUATION ----', text))
    if [int(match.group(1)) for match in markers] != list(
        range(1, protocol.EPOCHS + 1)
    ):
        raise ValueError('Training log does not contain ordered validation epochs 1..32.')

    records = []
    for index, marker in enumerate(markers):
        block_end = markers[index + 1].start() if index + 1 < len(markers) else len(text)
        block = text[marker.end():block_end]
        current = re.findall(
            r'Eval avg class IoU of prediction:\s*(%s)' % FLOAT_PATTERN,
            block,
        )
        best = re.findall(
            r'Best validation pixel IoU:\s*(%s) at epoch (\d+)' % FLOAT_PATTERN,
            block,
        )
        if len(current) != 1 or len(best) != 1:
            raise ValueError(
                'Epoch %d must contain one current-IoU and one cumulative-best record.'
                % (index + 1)
            )
        records.append((float(current[0]), float(best[0][0]), int(best[0][1])))

    tolerance = 1.1e-6  # train.py writes these three values with six decimals.
    current_values = []
    previous_best = -math.inf
    for epoch, (current, reported_best, reported_epoch) in enumerate(records, start=1):
        if not all(math.isfinite(value) for value in (current, reported_best)):
            raise ValueError('Training log contains non-finite validation IoU.')
        current_values.append(current)
        prefix_max = max(current_values)
        if abs(reported_best - prefix_max) > tolerance:
            raise ValueError('Logged cumulative best IoU disagrees with epoch IoU records.')
        if reported_epoch not in range(1, epoch + 1):
            raise ValueError('Logged cumulative best epoch is outside the observed prefix.')
        if abs(current_values[reported_epoch - 1] - reported_best) > tolerance:
            raise ValueError('Logged best epoch does not identify a maximum-IoU epoch.')
        if reported_best + tolerance < previous_best:
            raise ValueError('Logged cumulative best IoU decreases between epochs.')
        previous_best = reported_best
    final_current, final_best, final_epoch = records[-1]
    del final_current
    return final_epoch, final_best


def load_verified_rows(experiment, log_root, launcher=None):
    """Verify all seven run artifacts and return rows in registered order."""
    experiment = Path(experiment).resolve()
    log_root = Path(log_root).resolve()
    runner = _configure_launcher(experiment, log_root, launcher=launcher)
    jobs = protocol.read_manifest(experiment / 'manifest.tsv')
    val_names = _validation_names(runner)
    rows = []
    for job in jobs:
        done_path, elapsed = joint.validate_done(
            runner.status, job, require_exit_code=True,
        )
        # This binds the result to 32 complete validation records, the selected
        # best-IoU checkpoint, and the best-checkpoint evaluation payload.
        runner.validate_artifacts(job)

        metrics_path = runner.metrics_path(job).resolve()
        payload = json.loads(metrics_path.read_text(encoding='utf-8'))
        curve = _validated_curve(payload, val_names, job['run_id'] + ' metric')
        metrics = joint.recomputed_metrics(curve)

        checkpoint_path = Path(payload['checkpoint']).expanduser().resolve()
        checkpoint = load_checkpoint(checkpoint_path, map_location='cpu')
        selection = checkpoint['checkpoint_selection']
        validation = checkpoint['validation_metrics']
        best_epoch = int(selection['best_epoch'])
        best_val_iou = float(validation['iou'])
        if not math.isclose(
            best_val_iou, float(selection['best_value']),
            rel_tol=0.0, abs_tol=1e-12,
        ):
            raise ValueError('%s best validation IoU is inconsistent.' % job['run_id'])

        run_dir = runner.run_dir(job).resolve()
        training_log = run_dir / 'logs' / (protocol.MODEL + '.txt')
        logged_epoch, logged_iou = _logged_validation_selection(training_log)
        if logged_epoch != best_epoch or not math.isclose(
            logged_iou, best_val_iou, rel_tol=0.0, abs_tol=5.1e-7,
        ):
            raise ValueError(
                '%s checkpoint selection differs from the 32 validation records.'
                % job['run_id']
            )
        source_snapshot = run_dir / 'source_snapshot'
        joint.validate_source_snapshot(run_dir, complete_closure=True)
        expected_parameters = joint.EXPECTED_PARAMETERS[job['structure_variant']]
        if not math.isclose(
            float(payload['parameters_m']) * 1e6,
            expected_parameters,
            rel_tol=0.0,
            abs_tol=1e-6,
        ):
            raise ValueError(
                '%s parameter count must be %d.'
                % (job['run_id'], expected_parameters)
            )
        row = {
            'label': LABEL_BY_RUN_ID[job['run_id']],
            'run_id': job['run_id'],
            'variant': job['structure_variant'],
            'seed': int(job['seed']),
            'gpu': job['gpu'],
            'best_epoch': best_epoch,
            'best_val_iou': best_val_iou,
            'parameters': expected_parameters,
        }
        row.update(metrics)
        row.update({
            'elapsed_seconds': elapsed,
            'evaluation_seconds': float(payload['elapsed_seconds']),
            'metrics_path': str(metrics_path),
            'checkpoint_path': str(checkpoint_path),
            'training_log': str(training_log),
            'evaluation_log': str(run_dir / 'eval.txt'),
            'done_path': str(done_path),
            'source_snapshot': str(source_snapshot),
        })
        rows.append(row)

    joint.add_comparisons(rows)
    for row in rows:
        row['pareto'] = (
            'nondominated' if row['pareto_nondominated'] else 'dominated'
        )
    return rows


def _relative_path(path, report_root):
    return os.path.relpath(Path(path), start=Path(report_root))


def markdown_report(rows, report_root):
    lines = [
        '# BC-TPro corrected best-validation rerun results',
        '',
        '## 结果表',
        '',
        '| 模型 | 结构 | Seed | GPU | Best epoch | Best internal-val pixel IoU | Pd@0.5 (%) ↑ | Fa@0.5 (×10⁻⁵) ↓ | AUC27 ↑ | Pareto |',
        '|---|---|---:|---:|---:|---:|---:|---:|---:|---|',
    ]
    for row in rows:
        lines.append(
            '| {label} | `{variant}` | {seed} | {gpu} | {best_epoch} | '
            '{best_val_iou:.9f} | {pd:.6f} | {fa:.6f} | {auc27:.9f} | '
            '{pareto} |'.format(
                pd=row['pd_at_0_5'] * 100.0,
                fa=row['fa_at_0_5'] * 1e5,
                **row,
            )
        )
    frontier = [row['label'] for row in rows if row['pareto_nondominated']]
    lines += [
        '',
        '表内三目标 Pareto 非支配集：%s。' % '、'.join(frontier),
        '',
        '## 选择与比较规则',
        '',
    ]
    lines += ['- ' + note for note in PROTOCOL_NOTES]
    lines += ['', '## 相对参照', '']
    for label in ('B1', 'C1'):
        prefix = 'vs_' + label.lower() + '_'
        lines += [
            '### 相对 ' + label,
            '',
            '| 模型 | ΔPd (百分点) | ΔFa (绝对比例) | ΔAUC27 | Pareto关系 |',
            '|---|---:|---:|---:|---|',
        ]
        for row in rows:
            lines.append('| %s | %+.6f | %+.12g | %+.9f | %s |' % (
                row['label'], row[prefix + 'pd_pp'], row[prefix + 'fa_delta'],
                row[prefix + 'auc_delta'], row[prefix + 'relation'],
            ))
        lines.append('')
    lines += ['## 证据路径', '']
    for row in rows:
        links = []
        for label, key in (
            ('metrics', 'metrics_path'), ('checkpoint', 'checkpoint_path'),
            ('training log', 'training_log'), ('evaluation log', 'evaluation_log'),
            ('done marker', 'done_path'), ('source snapshot', 'source_snapshot'),
        ):
            links.append('[%s](%s)' % (
                label, _relative_path(row[key], report_root),
            ))
        lines.append('- %s：%s。' % (row['label'], '；'.join(links)))
    return '\n'.join(lines) + '\n'


def write_xlsx(rows, path):
    """Reuse the joint workbook, then add best-checkpoint selection fields."""
    if not joint.write_xlsx(rows, path):
        return False
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    summary = workbook['结果摘要']
    summary.insert_cols(5, amount=2)
    for source_column, target_column in ((7, 5), (7, 6)):
        source = summary.cell(1, source_column)
        target = summary.cell(1, target_column)
        target._style = copy(source._style)
        target.alignment = copy(source.alignment)
    summary.cell(1, 5).value = 'Best epoch'
    summary.cell(1, 6).value = 'Best internal-val pixel IoU ↑'
    for index, row in enumerate(rows, start=2):
        summary.cell(index, 5).value = row['best_epoch']
        summary.cell(index, 6).value = row['best_val_iou']
        summary.cell(index, 6).number_format = '0.000000000'
    summary.auto_filter.ref = summary.dimensions
    summary.column_dimensions['E'].width = 16
    summary.column_dimensions['F'].width = 34

    evidence = workbook['运行与产物']
    evidence.insert_cols(2, amount=2)
    for source_column, target_column in ((4, 2), (4, 3)):
        source = evidence.cell(1, source_column)
        target = evidence.cell(1, target_column)
        target._style = copy(source._style)
        target.alignment = copy(source.alignment)
    evidence.cell(1, 2).value = 'best_epoch'
    evidence.cell(1, 3).value = 'best_val_iou'
    for index, row in enumerate(rows, start=2):
        evidence.cell(index, 2).value = row['best_epoch']
        evidence.cell(index, 3).value = row['best_val_iou']
    evidence.auto_filter.ref = evidence.dimensions

    notes = workbook['协议与限制']
    notes.delete_rows(1, notes.max_row)
    notes.append(['说明'])
    for note in PROTOCOL_NOTES:
        notes.append([note])
    notes.column_dimensions['A'].width = 110
    notes.auto_filter.ref = notes.dimensions
    workbook.save(path)

    checked = load_workbook(path, read_only=True, data_only=True)
    if checked['结果摘要'].max_row != len(rows) + 1:
        raise ValueError('Workbook row count failed verification.')
    for index, row in enumerate(rows, start=2):
        sheet = checked['结果摘要']
        if sheet.cell(index, 5).value != row['best_epoch']:
            raise ValueError('Workbook best_epoch readback failed verification.')
        if not math.isclose(
            sheet.cell(index, 6).value, row['best_val_iou'],
            rel_tol=0.0, abs_tol=1e-14,
        ):
            raise ValueError('Workbook best_val_iou readback failed verification.')
    checked.close()
    return True


def write_reports(rows, experiment, xlsx=False):
    experiment = Path(experiment)
    if not experiment.is_dir():
        raise ValueError('Report directory must already exist.')
    with (experiment / 'results.csv').open(
        'w', newline='', encoding='utf-8-sig',
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    (experiment / 'RESULTS.md').write_text(
        markdown_report(rows, experiment), encoding='utf-8',
    )
    if xlsx:
        path = experiment / WORKBOOK_NAME
        if not write_xlsx(rows, path) or not path.is_file():
            raise RuntimeError(
                'Excel output requested but openpyxl is unavailable or output failed.'
            )


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--experiment-root', type=Path, default=DEFAULT_EXPERIMENT)
    parser.add_argument('--log-root', type=Path, default=DEFAULT_LOG_ROOT)
    parser.add_argument(
        '--xlsx', action='store_true',
        help='Also write the verified Excel workbook if openpyxl is installed.',
    )
    args = parser.parse_args(argv)
    experiment = args.experiment_root.resolve(strict=True)
    log_root = args.log_root.resolve(strict=True)
    rows = load_verified_rows(experiment, log_root)
    write_reports(rows, experiment, xlsx=args.xlsx)
    print(
        'VERIFIED: %d best-val runs; seed47; joint Pd/Fa/AUC Pareto; reports=%s'
        % (len(rows), experiment)
    )


if __name__ == '__main__':
    main()

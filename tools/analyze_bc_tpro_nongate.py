#!/usr/bin/env python3
"""Validate the single-seed nongated ablation and report joint Pd/Fa/AUC results.

Reuses the strict upstream checker, including checkpoint state, training and
evaluation provenance, and metrics recomputed from per-sequence integer counts.
No weighted score, metric-priority ordering, statistical claim, or candidate
lock is produced. Existing experiment artifacts are never changed.
"""

import argparse
import csv
import importlib
import json
import math
import sys
from pathlib import Path

import numpy as np

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from tools import analyze_bc_tpro_noise8_stage1 as legacy


DEFAULT_EXPERIMENT = REPO_ROOT / 'experiments/bc_tpro_nongate_noise8_seed47_2026-09-10'
REFERENCE_EXPERIMENT = REPO_ROOT / 'experiments/bc_tpro_stage1_noise8_upstream_2026-09-09'
NEW_VARIANTS = (
    ('ng1_ring_difference_seed47', 'center_ring_difference', 'NG1', '0'),
    ('ng2_bandpass_seed47', 'temporal_bandpass', 'NG2', '1'),
    ('ng3_spatial_smooth_seed47', 'center_spatial_smooth', 'NG3', '2'),
)
EXPECTED_PARAMETERS = {
    'none': 70913, 'temporal_control': 71233, 'center_multiscale': 71233,
    'center_ring': 71281,
    **{variant: 71233 for _, variant, _, _ in NEW_VARIANTS},
}
PROTOCOL_NOTES = (
    '状态：COMPLETED; ARTIFACTS_VERIFIED; SINGLE_SEED。所有指标经原始整数计数重算。',
    '数据集：NUDT-MIRSDT-Noise8.0_FJY；固定 train64/internal-val16；仅 seed47。',
    'FP32、scratch-only、固定 epoch32、T=40、batch4、crop128、Soft-IoU；无早停选模。',
    '联合目标为 Pd@0.5 越高、Fa@0.5 越低、官方 27 阈值 Pd-Fa AUC 越高。',
    'dominates：三项均非劣且至少一项严格改善；dominated：被参照支配；',
    'tradeoff：有改善也有退步；equal：三个未四舍五入数值完全相等。',
    'Pareto 非支配集表示没有被表内其他模型支配，不代表三项都优于基线或唯一最佳。',
    '差值均为当前行减参照；Fa 相对降幅为 (参照Fa-当前Fa)/参照Fa，正值表示减少。',
    '单seed没有训练随机性标准差、显著性结论或跨seed稳定性证据；没有自动候选锁。',
    '当前结果来自 internal-val16，不能直接与论文 official test20 数值比较。',
    '本轮不评测 official test20；该划分曾在项目历史实验使用，不是项目级从未见过的外部测试。',
    'NG1/NG2/NG3均为71,233参数；保留既有SiLU激活，没有新增可学习乘法门控。',
    'NG1固定中心减环形背景，NG2固定时间带通，NG3固定十字空间平滑；均用加法残差。',
    '等参数不意味着同噪声增益；NG3平滑同时改变噪声幅度与单像素目标响应。',
    'B1/C0/C1/C2复用已有seed47结果；C1/C2实际GPU为1/2，覆盖旧manifest的GPU0记录。',
    '行顺序为预登记 B1/C0/C1/C2/NG1/NG2/NG3；无加权分数，无单指标排序。',
)


def relation(candidate, reference):
    """Return the exact Pareto relation of finite, unrounded metric values."""
    difference = (
        candidate['pd_at_0_5'] - reference['pd_at_0_5'],
        reference['fa_at_0_5'] - candidate['fa_at_0_5'],
        candidate['auc27'] - reference['auc27'],
    )
    if not all(math.isfinite(value) for value in difference):
        raise ValueError('Pareto comparison requires finite metrics.')
    if all(value == 0 for value in difference):
        return 'equal'
    if all(value >= 0 for value in difference):
        return 'dominates'
    if all(value <= 0 for value in difference):
        return 'dominated'
    return 'tradeoff'


def add_comparisons(rows):
    references = {row['label']: row for row in rows}
    for label in ('B1', 'C1'):
        reference = references[label]
        for row in rows:
            prefix = 'vs_' + label.lower() + '_'
            row[prefix + 'pd_pp'] = 100 * (row['pd_at_0_5'] - reference['pd_at_0_5'])
            row[prefix + 'fa_delta'] = row['fa_at_0_5'] - reference['fa_at_0_5']
            row[prefix + 'fa_relative_reduction'] = (
                -row[prefix + 'fa_delta'] / reference['fa_at_0_5']
                if reference['fa_at_0_5'] > 0 else None
            )
            row[prefix + 'auc_delta'] = row['auc27'] - reference['auc27']
            row[prefix + 'relation'] = relation(row, reference)
    for row in rows:
        row['pareto_nondominated'] = not any(
            relation(other, row) == 'dominates' for other in rows
        )
    return rows


def validate_done(status_root, job, require_exit_code=False):
    """Accept historical key=value and structured completion markers."""
    status_root = Path(status_root)
    for suffix in ('running', 'failed'):
        if (status_root / (job['run_id'] + '.' + suffix)).exists():
            raise ValueError('Conflicting %s marker for %s.' % (suffix, job['run_id']))
    path = status_root / (job['run_id'] + '.done')
    text = path.read_text(encoding='utf-8').strip()
    if text.startswith('{'):
        values = json.loads(text)
        if not isinstance(values, dict):
            raise ValueError('Completion marker must be a mapping: %s' % path)
    else:
        values = {}
        for line in text.splitlines():
            key, separator, value = line.partition('=')
            if not separator or not key or key in values:
                raise ValueError('Malformed completion marker: %s' % path)
            values[key] = value
    for key in ('run_id', 'wave', 'gpu', 'seed'):
        if str(values.get(key)) != str(job[key]):
            raise ValueError('Completion %s mismatch: %s' % (key, path))
    if require_exit_code and 'exit_code' not in values:
        raise ValueError('Completion exit_code missing: %s' % path)
    if str(values.get('exit_code', 0)) != '0':
        raise ValueError('Completion exit_code is not zero: %s' % path)
    if not values.get('started_at') or not values.get('finished_at'):
        raise ValueError('Completion timestamps missing: %s' % path)
    elapsed = float(values.get('elapsed_seconds', float('nan')))
    if not math.isfinite(elapsed) or elapsed <= 0:
        raise ValueError('Invalid completion elapsed_seconds: %s' % path)
    return path, elapsed


def validate_source_snapshot(run_dir, complete_closure=False):
    """Verify retained copies without comparing historical code to today's code."""
    from tools.run_bc_tpro_nongate import (
        EXTRA_SNAPSHOT_FILES, PACKAGES, SNAPSHOT_FILES,
    )
    snapshot = run_dir / 'source_snapshot'
    required = list(SNAPSHOT_FILES)
    if complete_closure:
        required += list(EXTRA_SNAPSHOT_FILES)
        required += [package + '/__init__.py' for package in PACKAGES]
    for relative in required:
        if not (snapshot / relative).is_file():
            raise ValueError('Missing source snapshot: %s' % (snapshot / relative))
    for relative in ('networks/models/DeepPro-Plus_BCTPro.py',
                     'networks/layers/bc_tpro_adapter.py',
                     'networks/losses/segmentation_losses.py'):
        root_copy = run_dir / Path(relative).name
        if not root_copy.is_file() or root_copy.read_bytes() != (snapshot / relative).read_bytes():
            raise ValueError('Root and nested source snapshots differ: %s' % root_copy)


def load_new_manifest(experiment):
    with (experiment / 'manifest.tsv').open(newline='', encoding='utf-8') as handle:
        reader = csv.DictReader(handle, delimiter='\t')
        columns = ['run_id', 'wave', 'model', 'structure_variant', 'seed', 'gpu', 'log_dir']
        if reader.fieldnames != columns:
            raise ValueError('New manifest columns differ from the registered design.')
        jobs = list(reader)
    expected = []
    for run_id, variant, label, gpu in NEW_VARIANTS:
        expected.append(dict(
            run_id=run_id, wave='1', model=legacy.MODEL,
            structure_variant=variant, seed='47', gpu=gpu,
            log_dir=('2026-09-10/%s__Upstream8fa1a68-FP32-SoftIoU-%s_seed47_E32'
                     % (legacy.DATASET, label)),
        ))
    if jobs != expected:
        raise ValueError('New manifest differs from the registered three-run design.')
    return jobs


def recomputed_metrics(curve):
    true_counts = curve['true_counts'].sum(axis=0)
    total_counts = curve['target_counts'].sum(axis=0)
    false_counts = curve['false_counts'].sum(axis=0)
    pixels = float(curve['pixels'].sum())
    pd = true_counts / total_counts
    fa = false_counts / pixels
    indices = [int(np.flatnonzero(curve['thresholds'] == threshold)[0])
               for threshold in legacy.common.PAPER_THRESHOLDS]
    half = curve['threshold_half_index']
    return dict(
        pd_at_0_5=float(pd[half]), fa_at_0_5=float(fa[half]),
        auc27=float(abs(np.trapz(pd[indices], fa[indices]))),
        true_targets=int(true_counts[half]), total_targets=int(total_counts[half]),
        false_pixels=int(false_counts[half]), pixel_count=int(pixels),
    )


def load_verified_rows(experiment, reference_experiment, log_root):
    new_jobs = load_new_manifest(experiment)
    dataset_root = (REPO_ROOT.parent / 'datasets' / legacy.DATASET).resolve()
    train_list, val_list, val_names = legacy.validate_splits(reference_experiment, dataset_root)
    reference_jobs = [dict(job) for job in legacy.expected_manifest_rows(legacy.UPSTREAM_PROFILE)
                      if job['seed'] == '47']
    for job in reference_jobs:
        job['gpu'] = {'none': '0', 'temporal_control': '0',
                      'center_multiscale': '1', 'center_ring': '2'}[job['structure_variant']]
    entries = [(job, reference_experiment, label) for job, label in
               zip(reference_jobs, ('B1', 'C0', 'C1', 'C2'))]
    entries += [(job, experiment, variant[2]) for job, variant in zip(new_jobs, NEW_VARIANTS)]
    checkpoint_cache, training_cache, rows = {}, {}, []
    model_module = importlib.import_module('networks.models.DeepPro-Plus_BCTPro')
    for job, owner, label in entries:
        status_root = log_root / '_queues' / owner.name / 'status'
        done_path, elapsed = validate_done(status_root, job, require_exit_code=owner == experiment)
        validate_source_snapshot(log_root / job['log_dir'], complete_closure=owner == experiment)
        metrics_path = owner / 'metrics' / (job['run_id'] + legacy.METRICS_SUFFIX)
        payload = json.loads(metrics_path.read_text(encoding='utf-8'))
        curve = legacy.validate_one_payload(
            payload, job, metrics_path, owner, log_root, dataset_root,
            train_list, val_list, val_names, checkpoint_cache, training_cache,
            legacy.UPSTREAM_PROFILE,
        )
        model = model_module.detector(1, 40, 40, structure_variant=job['structure_variant'],
                                     structure_bottleneck_channels=8, eval_chunk_rows=32)
        parameters = sum(parameter.numel() for parameter in model.parameters())
        expected = EXPECTED_PARAMETERS[job['structure_variant']]
        if parameters != expected or not math.isclose(
            float(payload['parameters_m']) * 1e6, expected, rel_tol=0, abs_tol=1e-6,
        ):
            raise ValueError('%s parameter count must be %d.' % (job['run_id'], expected))
        row = dict(label=label, run_id=job['run_id'], variant=job['structure_variant'],
                   seed=47, gpu=job['gpu'], parameters=parameters)
        row.update(recomputed_metrics(curve))
        row.update(elapsed_seconds=elapsed, evaluation_seconds=float(payload['elapsed_seconds']),
                   metrics_path=str(metrics_path), checkpoint_path=payload['checkpoint'],
                   done_path=str(done_path),
                   training_log=str(log_root / job['log_dir'] / 'logs' / (legacy.MODEL + '.txt')),
                   evaluation_log=str(log_root / job['log_dir'] / 'eval_epoch-32.txt'),
                   source_snapshot=str(log_root / job['log_dir'] / 'source_snapshot'))
        rows.append(row)
    return add_comparisons(rows)


def markdown_report(rows):
    lines = ['# BC-TPro 无门控结构实验结果', '', '## Material Passport', '',
             '- Origin workflow: academic-research-suite / experiment-agent / analyze',
             '- Verification: COMPLETED; ARTIFACTS_VERIFIED; SINGLE_SEED', '',
             '| 模型 | Seed | GPU | 参数量 | Pd@0.5 (%) ↑ | Fa@0.5 (×10⁻⁵) ↓ | AUC27 ↑ | Pareto非支配 |',
             '|---|---:|---:|---:|---:|---:|---:|---|']
    for row in rows:
        lines.append('| {label} | {seed} | {gpu} | {parameters} | {pd:.6f} | {fa:.6f} | {auc27:.9f} | {frontier} |'.format(
            pd=row['pd_at_0_5'] * 100, fa=row['fa_at_0_5'] * 1e5,
            frontier='是' if row['pareto_nondominated'] else '否', **row))
    frontier = [row['label'] for row in rows if row['pareto_nondominated']]
    lines += ['', '表内 Pareto 非支配集：' + '、'.join(frontier) + '。']
    for label in ('B1', 'C1'):
        prefix = 'vs_' + label.lower() + '_'
        lines += ['', '## 相对 ' + label, '',
                  '| 模型 | ΔPd (百分点) | ΔFa (绝对比例) | Fa相对降幅 (%) | ΔAUC27 | 关系 |',
                  '|---|---:|---:|---:|---:|---|']
        for row in rows:
            reduction = row[prefix + 'fa_relative_reduction']
            lines.append('| %s | %+.6f | %+.12g | %s | %+.9f | %s |' % (
                row['label'], row[prefix + 'pd_pp'], row[prefix + 'fa_delta'],
                '%+.6f' % (reduction * 100) if reduction is not None else '未定义',
                row[prefix + 'auc_delta'], row[prefix + 'relation']))
    lines += ['', '## 协议与限制', ''] + ['- ' + note for note in PROTOCOL_NOTES]
    lines += ['', '## 原始证据', '']
    for row in rows:
        lines.append('- %s：[指标](%s)、[checkpoint](%s)、[训练日志](%s)、[评测日志](%s)、[完成标记](%s)。' % (
            row['label'], row['metrics_path'], row['checkpoint_path'], row['training_log'],
            row['evaluation_log'], row['done_path']))
    return '\n'.join(lines) + '\n'


def write_xlsx(rows, path):
    """Optional workbook; dependency absence does not block CSV/Markdown."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print('openpyxl unavailable; CSV and Markdown remain complete.', file=sys.stderr)
        return False
    workbook = Workbook()
    summary = workbook.active
    summary.title = '结果摘要'
    summary.append(['模型', '结构', 'Seed', 'GPU', '参数量', 'Pd@0.5 (%) ↑',
                    'Fa@0.5 (×10⁻⁵) ↓', 'AUC27 ↑', 'Pareto非支配'])
    for row in rows:
        summary.append([row['label'], row['variant'], row['seed'], row['gpu'], row['parameters'],
                        row['pd_at_0_5'] * 100, row['fa_at_0_5'] * 1e5, row['auc27'],
                        '是' if row['pareto_nondominated'] else '否'])
    for label in ('B1', 'C1'):
        sheet = workbook.create_sheet('相对' + label)
        sheet.append(['模型', 'ΔPd (百分点)', 'ΔFa (绝对比例)', 'Fa相对降幅', 'ΔAUC27', '关系'])
        prefix = 'vs_' + label.lower() + '_'
        for row in rows:
            sheet.append([row['label']] + [row[prefix + field] for field in
                         ('pd_pp', 'fa_delta', 'fa_relative_reduction', 'auc_delta', 'relation')])
            sheet.cell(sheet.max_row, 4).number_format = '0.0000%'
            sheet.cell(sheet.max_row, 3).number_format = '0.000000E+00'
    evidence = workbook.create_sheet('运行与产物')
    fields = ['label', 'elapsed_seconds', 'evaluation_seconds', 'metrics_path', 'checkpoint_path',
              'training_log', 'evaluation_log', 'done_path', 'source_snapshot']
    evidence.append(fields)
    for row in rows:
        evidence.append([row[field] for field in fields])
        for column in range(4, len(fields) + 1):
            cell = evidence.cell(evidence.max_row, column)
            cell.hyperlink = cell.value
    notes = workbook.create_sheet('协议与限制')
    notes.append(['说明'])
    for note in PROTOCOL_NOTES:
        notes.append([note])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='183B56')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 32
        for column in sheet.columns:
            width = min(65, max(16, max(len(str(cell.value or '')) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    notes.column_dimensions['A'].width = 110
    for cell in notes['A']:
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for row in range(2, summary.max_row + 1):
        for column in (6, 7, 8):
            summary.cell(row, column).number_format = '0.000000'
    workbook.save(path)
    checked = load_workbook(path, read_only=True, data_only=True)
    if checked['结果摘要'].max_row != len(rows) + 1:
        raise ValueError('Workbook row count failed verification.')
    for index, row in enumerate(rows, start=2):
        if not math.isclose(checked['结果摘要'].cell(index, 8).value, row['auc27'],
                            rel_tol=0, abs_tol=1e-14):
            raise ValueError('Workbook AUC readback failed verification.')
    checked.close()
    return True


def write_reports(rows, experiment, xlsx=False):
    experiment = Path(experiment)
    if not experiment.is_dir():
        raise ValueError('Report directory must already exist.')
    with (experiment / 'results.csv').open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    (experiment / 'RESULTS.md').write_text(markdown_report(rows), encoding='utf-8')
    if xlsx:
        write_xlsx(rows, experiment / 'NG_EXPERIMENT_RESULTS_2026-09-10.xlsx')


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--experiment-root', type=Path, default=DEFAULT_EXPERIMENT)
    parser.add_argument('--reference-experiment', type=Path, default=REFERENCE_EXPERIMENT)
    parser.add_argument('--log-root', type=Path, default=REPO_ROOT / 'log/sem_seg')
    parser.add_argument('--xlsx', action='store_true', help='Also write Excel if openpyxl is installed.')
    args = parser.parse_args(argv)
    experiment = args.experiment_root.resolve(strict=True)
    reference = args.reference_experiment.resolve(strict=True)
    if experiment == reference:
        raise ValueError('New reports must not overwrite the historical reference experiment.')
    rows = load_verified_rows(experiment, reference, args.log_root.resolve(strict=True))
    write_reports(rows, experiment, xlsx=args.xlsx)
    print('VERIFIED: %d runs; single seed47; joint Pd/Fa/AUC; reports=%s' % (len(rows), experiment))


if __name__ == '__main__':
    main()

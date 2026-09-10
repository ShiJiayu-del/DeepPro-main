"""CPU-only tests for joint metric comparisons and report integrity."""

import csv
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

import numpy as np

from tools import analyze_bc_tpro_nongate as analysis


def metric(pd, fa, auc):
    return dict(pd_at_0_5=pd, fa_at_0_5=fa, auc27=auc)


def dummy_rows():
    rows = []
    for index, label in enumerate(('B1', 'C0', 'C1', 'C2', 'NG1', 'NG2', 'NG3')):
        row = dict(label=label, run_id=label.lower() + '_seed47', variant=label,
                   seed=47, gpu=str(index % 3), parameters=71233,
                   elapsed_seconds=100, evaluation_seconds=10,
                   metrics_path='/evidence/' + label + '.json', checkpoint_path='/evidence/' + label + '.pth',
                   training_log='/evidence/train.txt', evaluation_log='/evidence/eval.txt',
                   done_path='/evidence/' + label + '.done', source_snapshot='/evidence/sources')
        row.update(metric(0.7 + index * 0.01, 4e-5 - index * 1e-6, 0.9 + index * 0.005))
        rows.append(row)
    return analysis.add_comparisons(rows)


class JointMetricTests(unittest.TestCase):
    def test_all_pareto_relations_include_ties(self):
        reference = metric(0.8, 3e-5, 0.94)
        examples = (
            (metric(0.81, 3e-5, 0.94), 'dominates'),
            (metric(0.8, 2e-5, 0.94), 'dominates'),
            (metric(0.8, 3e-5, 0.93), 'dominated'),
            (metric(0.79, 2e-5, 0.95), 'tradeoff'),
            (metric(0.8, 3e-5, 0.94), 'equal'),
        )
        for candidate, expected in examples:
            self.assertEqual(analysis.relation(candidate, reference), expected)
        with self.assertRaisesRegex(ValueError, 'finite'):
            analysis.relation(metric(float('nan'), 3e-5, 0.94), reference)

    def test_tradeoff_is_not_mistaken_for_dominance(self):
        rows = [dict(label='B1', **metric(0.8, 3e-5, 0.94)),
                dict(label='C1', **metric(0.79, 2e-5, 0.95)),
                dict(label='NG1', **metric(0.8, 2e-5, 0.96))]
        analysis.add_comparisons(rows)
        self.assertEqual(rows[1]['vs_b1_relation'], 'tradeoff')
        self.assertEqual([row['label'] for row in rows if row['pareto_nondominated']], ['NG1'])
        self.assertAlmostEqual(rows[1]['vs_b1_pd_pp'], -1)
        self.assertAlmostEqual(rows[1]['vs_b1_fa_delta'], -1e-5)
        self.assertAlmostEqual(rows[1]['vs_b1_fa_relative_reduction'], 1 / 3)
        self.assertEqual(rows[0]['vs_b1_relation'], 'equal')

    def test_zero_reference_fa_is_explicitly_undefined(self):
        rows = [dict(label='B1', **metric(0.8, 0, 0.94)),
                dict(label='C1', **metric(0.8, 1e-5, 0.95))]
        analysis.add_comparisons(rows)
        self.assertIsNone(rows[1]['vs_b1_fa_relative_reduction'])

    def test_auc_recomputed_on_paper_threshold_subset(self):
        thresholds = analysis.legacy.common.EXPECTED_THRESHOLD_GRID
        count = len(thresholds)
        curve = dict(thresholds=thresholds,
                     true_counts=np.array([np.linspace(10, 0, count)]),
                     target_counts=np.full((1, count), 10),
                     false_counts=np.array([np.linspace(100, 0, count)]),
                     pixels=np.array([100]),
                     threshold_half_index=int(np.flatnonzero(thresholds == 0.5)[0]))
        values = analysis.recomputed_metrics(curve)
        self.assertAlmostEqual(values['auc27'], 0.5)


class ReportTests(unittest.TestCase):
    def test_completed_markers_accept_historical_and_json_formats(self):
        job = dict(run_id='ng1_seed47', wave='1', gpu='0', seed='47')
        values = dict(job, started_at='2026-09-10T00:00:00Z',
                      finished_at='2026-09-10T01:00:00Z', elapsed_seconds=3600)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / 'ng1_seed47.done'
            for text in ('\n'.join('%s=%s' % item for item in values.items()), json.dumps(values)):
                path.write_text(text, encoding='utf-8')
                self.assertEqual(analysis.validate_done(root, job)[1], 3600)
            with self.assertRaisesRegex(ValueError, 'exit_code missing'):
                analysis.validate_done(root, job, require_exit_code=True)
            path.write_text(json.dumps(dict(values, exit_code=0)), encoding='utf-8')
            self.assertEqual(analysis.validate_done(root, job, require_exit_code=True)[1], 3600)
            wrong = dict(values, gpu='3')
            path.write_text(json.dumps(wrong), encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'gpu mismatch'):
                analysis.validate_done(root, job)
            path.write_text(json.dumps(dict(values, exit_code=1)), encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'exit_code'):
                analysis.validate_done(root, job)

    def test_source_snapshot_rejects_missing_or_changed_root_adapter(self):
        from tools.run_bc_tpro_nongate import SNAPSHOT_FILES
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaisesRegex(ValueError, 'Missing source snapshot'):
                analysis.validate_source_snapshot(root)
            for relative in SNAPSHOT_FILES:
                target = root / 'source_snapshot' / relative
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text(relative, encoding='utf-8')
                if relative in ('networks/models/DeepPro-Plus_BCTPro.py',
                                'networks/layers/bc_tpro_adapter.py',
                                'networks/losses/segmentation_losses.py'):
                    (root / Path(relative).name).write_text(relative, encoding='utf-8')
            analysis.validate_source_snapshot(root)
            with self.assertRaisesRegex(ValueError, 'Missing source snapshot'):
                analysis.validate_source_snapshot(root, complete_closure=True)
            (root / 'bc_tpro_adapter.py').write_text('changed', encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'snapshots differ'):
                analysis.validate_source_snapshot(root)

    def test_csv_and_markdown_preserve_registered_order_and_precision(self):
        rows = dummy_rows()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            analysis.write_reports(rows, root)
            with (root / 'results.csv').open(encoding='utf-8-sig', newline='') as handle:
                restored = list(csv.DictReader(handle))
            self.assertEqual([row['label'] for row in restored], [row['label'] for row in rows])
            self.assertEqual(float(restored[-1]['auc27']), rows[-1]['auc27'])
            report = (root / 'RESULTS.md').read_text(encoding='utf-8')
            for text in ('SINGLE_SEED', 'internal-val16', '项目历史', 'SiLU', '相对 B1', '相对 C1'):
                self.assertIn(text, report)
            self.assertEqual(len(list(root.iterdir())), 2)

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'openpyxl optional')
    def test_optional_workbook_values_and_seven_rows(self):
        from openpyxl import load_workbook
        rows = dummy_rows()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'results.xlsx'
            self.assertTrue(analysis.write_xlsx(rows, path))
            workbook = load_workbook(path, data_only=True)
            self.assertEqual(workbook['结果摘要'].max_row, 8)
            self.assertAlmostEqual(workbook['相对B1']['B8'].value, rows[-1]['vs_b1_pd_pp'])
            self.assertEqual(workbook['相对C1']['F4'].value, 'equal')
            workbook.close()


if __name__ == '__main__':
    unittest.main()

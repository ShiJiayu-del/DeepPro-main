"""CPU-only tests for the corrected best-validation result analyzer."""

import csv
import importlib.util
import json
import math
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import numpy as np
import torch

from tools import analyze_bc_tpro_bestval as analysis
from tools.run_bc_tpro_nongate import EXTRA_SNAPSHOT_FILES, PACKAGES, SNAPSHOT_FILES


def aggregate_metrics(thresholds, true_counts, target_counts, false_counts, pixels):
    true_totals = true_counts.sum(axis=0)
    target_totals = target_counts.sum(axis=0)
    false_totals = false_counts.sum(axis=0)
    pd = true_totals / target_totals
    fa = false_totals / float(pixels.sum())
    half = int(np.flatnonzero(thresholds == 0.5)[0])
    paper = [
        int(np.flatnonzero(thresholds == value)[0])
        for value in analysis.joint.legacy.common.PAPER_THRESHOLDS
    ]
    return {
        'true_targets': float(true_totals[half]),
        'total_targets': float(target_totals[half]),
        'false_pixels': float(false_totals[half]),
        'pixel_count': float(pixels.sum()),
        'pd': float(pd[half]),
        'fa': float(fa[half]),
        'pd_percent': float(pd[half] * 100.0),
        'fa_x1e5': float(fa[half] * 1e5),
        'auc': float(abs(np.trapz(pd[paper], fa[paper]))),
        'auc_dense_grid': float(abs(np.trapz(pd, fa))),
    }


class BestValidationFixture:
    CURVE_SCALES = (
        (80, 1000),
        (70, 1000),
        (90, 900),
        (90, 800),
        (95, 900),
        (95, 800),
        (60, 1100),
    )

    def __init__(self, root):
        self.root = Path(root) / 'repo'
        self.root.mkdir(parents=True)
        self.runner = analysis.protocol.Launcher(self.root)
        self.experiment = self.runner.experiment
        self.log_root = self.root / 'log' / 'sem_seg'
        self.jobs = analysis.protocol.expected_manifest_rows()
        self.experiment.mkdir(parents=True)
        with (self.experiment / 'manifest.tsv').open(
            'w', newline='', encoding='utf-8',
        ) as handle:
            writer = csv.DictWriter(
                handle, analysis.protocol.COLUMNS, delimiter='\t',
            )
            writer.writeheader()
            writer.writerows(self.jobs)
        split = self.runner.split('val_sequences.txt')
        split.parent.mkdir(parents=True)
        split.write_text(
            '\n'.join(analysis.joint.legacy.VAL_NAMES) + '\n', encoding='utf-8',
        )
        self.runner.status.mkdir(parents=True)
        (self.experiment / 'metrics').mkdir()
        for index, job in enumerate(self.jobs):
            self.create_run(job, index, *self.CURVE_SCALES[index])

    def payload(self, job, checkpoint_path, best_epoch, pd_scale, fa_scale):
        thresholds = np.asarray(
            analysis.joint.legacy.common.EXPECTED_THRESHOLD_GRID,
            dtype=np.float64,
        )
        sequence_count = len(analysis.joint.legacy.VAL_NAMES)
        target_counts = np.full((sequence_count, thresholds.size), 100, dtype=np.int64)
        true_row = np.floor(pd_scale * (1.0 - thresholds)).astype(np.int64)
        false_row = np.floor(fa_scale * (1.0 - thresholds)).astype(np.int64)
        true_counts = np.tile(true_row, (sequence_count, 1))
        false_counts = np.tile(false_row, (sequence_count, 1))
        pixels = np.full(sequence_count, 10000, dtype=np.int64)
        all_metrics = aggregate_metrics(
            thresholds, true_counts, target_counts, false_counts, pixels,
        )
        return {
            'schema_version': 2,
            'dataset': analysis.protocol.DATASET,
            'model': analysis.protocol.MODEL,
            'checkpoint': str(checkpoint_path),
            'checkpoint_epoch': best_epoch,
            'sequence_length': 40,
            'sequence_count': sequence_count,
            'evaluation_windows': analysis.protocol.VAL_EVALUATION_WINDOWS,
            'inference_amp': False,
            'operating_threshold': 0.5,
            'paper_thresholds': analysis.joint.legacy.common.PAPER_THRESHOLDS.tolist(),
            'elapsed_seconds': 2.5,
            'parameters_m': (
                analysis.joint.EXPECTED_PARAMETERS[job['structure_variant']] / 1e6
            ),
            'cuda_peak_allocated_gib': 1.0,
            'cuda_peak_reserved_gib': 1.25,
            'all': all_metrics,
            'curve_counts': {
                'sequence_names': list(analysis.joint.legacy.VAL_NAMES),
                'thresholds': thresholds.tolist(),
                'false_pixels_by_sequence': false_counts.tolist(),
                'true_targets_by_sequence': true_counts.tolist(),
                'total_targets_by_sequence': target_counts.tolist(),
                'pixel_count_by_sequence': pixels.tolist(),
            },
        }

    def create_run(self, job, index, pd_scale, fa_scale):
        best_epoch = index + 3
        best_iou = 0.55 + index * 0.01
        run_dir = self.runner.run_dir(job)
        training_log = run_dir / 'logs' / (analysis.protocol.MODEL + '.txt')
        training_log.parent.mkdir(parents=True)
        running_best = -math.inf
        running_epoch = 0
        records = []
        for epoch in range(1, analysis.protocol.EPOCHS + 1):
            current_iou = best_iou if epoch == best_epoch else best_iou - 0.1
            if current_iou >= running_best:
                running_best = current_iou
                running_epoch = epoch
            records.append(
                (
                    '---- EPOCH %03d EVALUATION ----\n'
                    'Eval mean loss: 0.400000\n'
                    'Eval avg class IoU of prediction: %.6f\n'
                    'Eval pixel precision: 0.700000\n'
                    'Eval pixel recall: 0.600000\n'
                    'Eval pixel F1: 0.646154\n'
                    'Best validation pixel IoU: %.6f at epoch %d\n'
                ) % (epoch, current_iou, running_best, running_epoch)
            )
        training_log.write_text(
            ''.join(records),
            encoding='utf-8',
        )
        checkpoint_path = (run_dir / 'checkpoints' / 'best_model.pth').resolve()
        checkpoint_path.parent.mkdir()
        torch.save(
            {
                'epoch': best_epoch - 1,
                'model_name': analysis.protocol.MODEL,
                'model_config': {'structure_variant': job['structure_variant']},
                'checkpoint_selection': {
                    'metric': 'eval_iou', 'mode': 'max',
                    'overlap_policy': 'official_window',
                    'best_value': best_iou, 'best_epoch': best_epoch,
                },
                'validation_metrics': {
                    'epoch': best_epoch, 'loss': 1.0 - best_iou,
                    'overlap_policy': 'official_window',
                    'iou': best_iou, 'precision': 0.7, 'recall': 0.6,
                    'f1': 0.646153846,
                },
            },
            checkpoint_path,
        )
        metrics_path = self.runner.metrics_path(job).resolve()
        payload = self.payload(
            job, checkpoint_path, best_epoch, pd_scale, fa_scale,
        )
        metrics_path.write_text(json.dumps(payload) + '\n', encoding='utf-8')
        (run_dir / 'eval.txt').write_text(
            'Paper-aligned metrics saved to %s.\n' % metrics_path,
            encoding='utf-8',
        )
        snapshot = run_dir / 'source_snapshot'
        for relative in SNAPSHOT_FILES + EXTRA_SNAPSHOT_FILES:
            target = snapshot / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(relative, encoding='utf-8')
        for package in PACKAGES:
            target = snapshot / package / '__init__.py'
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.exists():
                target.write_text('', encoding='utf-8')
        for relative in (
            'networks/models/DeepPro-Plus_BCTPro.py',
            'networks/layers/bc_tpro_adapter.py',
            'networks/losses/segmentation_losses.py',
        ):
            (run_dir / Path(relative).name).write_bytes((snapshot / relative).read_bytes())
        done = {
            'run_id': job['run_id'], 'wave': job['wave'], 'gpu': job['gpu'],
            'seed': job['seed'], 'started_at': '2026-09-11T00:00:00+08:00',
            'finished_at': '2026-09-11T01:00:00+08:00', 'exit_code': 0,
            'elapsed_seconds': 3600,
        }
        (self.runner.status / (job['run_id'] + '.done')).write_text(
            ''.join('%s=%s\n' % item for item in done.items()), encoding='utf-8',
        )


class BestValidationAnalyzerTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.fixture = BestValidationFixture(self.temporary.name)

    def load(self):
        return analysis.load_verified_rows(
            self.fixture.experiment,
            self.fixture.log_root,
            launcher=self.fixture.runner,
        )

    def test_seven_rows_are_verified_recomputed_and_pareto_classified(self):
        with mock.patch.object(
            self.fixture.runner,
            'validate_artifacts',
            wraps=self.fixture.runner.validate_artifacts,
        ) as validate:
            rows = self.load()
        self.assertEqual(validate.call_count, 7)
        self.assertEqual(
            [row['label'] for row in rows],
            ['B1', 'C0', 'C1', 'C2', 'NG1', 'NG2', 'NG3'],
        )
        self.assertEqual([row['best_epoch'] for row in rows], list(range(3, 10)))
        self.assertAlmostEqual(rows[0]['best_val_iou'], 0.55)
        self.assertAlmostEqual(rows[0]['pd_at_0_5'], 0.4)
        self.assertAlmostEqual(rows[0]['fa_at_0_5'], 0.05)
        self.assertEqual(rows[0]['true_targets'], 640)
        self.assertEqual(rows[0]['false_pixels'], 8000)
        self.assertEqual(
            [row['label'] for row in rows if row['pareto_nondominated']],
            ['NG1', 'NG2'],
        )
        self.assertEqual(rows[4]['vs_c1_relation'], 'dominates')
        self.assertEqual(rows[5]['vs_c1_relation'], 'tradeoff')
        self.assertEqual(rows[1]['pareto'], 'dominated')

    def test_done_marker_and_integer_curve_are_mandatory(self):
        job = self.fixture.jobs[0]
        done = self.fixture.runner.status / (job['run_id'] + '.done')
        done.rename(done.with_suffix('.failed'))
        with self.assertRaisesRegex(ValueError, 'Conflicting failed marker'):
            self.load()

        done.with_suffix('.failed').rename(done)
        metrics_path = self.fixture.runner.metrics_path(job)
        payload = json.loads(metrics_path.read_text(encoding='utf-8'))
        payload['curve_counts']['true_targets_by_sequence'][0][0] = 1.5
        metrics_path.write_text(json.dumps(payload), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'integer counts'):
            self.load()

    def test_checkpoint_selection_must_match_epoch_validation_log(self):
        job = self.fixture.jobs[0]
        training_log = (
            self.fixture.runner.run_dir(job) / 'logs' / (analysis.protocol.MODEL + '.txt')
        )
        text = training_log.read_text(encoding='utf-8')
        text = text.replace(
            'Best validation pixel IoU: 0.550000 at epoch 3\n',
            'Best validation pixel IoU: 0.990000 at epoch 32\n',
            1,
        )
        training_log.write_text(text, encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'cumulative best IoU'):
            self.load()

    def test_parameter_count_and_complete_snapshot_are_verified(self):
        job = self.fixture.jobs[0]
        metrics_path = self.fixture.runner.metrics_path(job)
        payload = json.loads(metrics_path.read_text(encoding='utf-8'))
        payload['parameters_m'] = 123.456
        metrics_path.write_text(json.dumps(payload), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'parameter count'):
            self.load()

        self.fixture = BestValidationFixture(
            Path(self.temporary.name) / 'snapshot-case'
        )
        missing = (
            self.fixture.runner.run_dir(self.fixture.jobs[0])
            / 'source_snapshot' / 'train.py'
        )
        missing.unlink()
        with self.assertRaisesRegex(ValueError, 'Missing source snapshot'):
            self.load()

    def test_csv_and_markdown_include_selection_metrics_and_evidence(self):
        rows = self.load()
        analysis.write_reports(rows, self.fixture.experiment)
        with (self.fixture.experiment / 'results.csv').open(
            newline='', encoding='utf-8-sig',
        ) as handle:
            restored = list(csv.DictReader(handle))
        self.assertEqual(len(restored), 7)
        for field in (
            'label', 'variant', 'seed', 'gpu', 'best_epoch', 'best_val_iou',
            'pd_at_0_5', 'fa_at_0_5', 'auc27', 'pareto', 'metrics_path',
            'checkpoint_path', 'training_log', 'evaluation_log', 'done_path',
        ):
            self.assertIn(field, restored[0])
        report = (self.fixture.experiment / 'RESULTS.md').read_text(encoding='utf-8')
        for text in (
            'Best internal-val pixel IoU', 'best_model.pth',
            '不存在 AUC 优先规则', '三目标 Pareto', '证据路径',
        ):
            self.assertIn(text, report)

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'openpyxl optional')
    def test_xlsx_has_seven_rows_and_best_selection_columns(self):
        from openpyxl import load_workbook

        rows = self.load()
        path = self.fixture.experiment / analysis.WORKBOOK_NAME
        self.assertTrue(analysis.write_xlsx(rows, path))
        workbook = load_workbook(path, read_only=True, data_only=True)
        summary = workbook['结果摘要']
        self.assertEqual(summary.max_row, 8)
        self.assertEqual(summary.cell(1, 5).value, 'Best epoch')
        self.assertEqual(
            summary.cell(1, 6).value, 'Best internal-val pixel IoU ↑',
        )
        self.assertEqual(summary.cell(8, 5).value, rows[-1]['best_epoch'])
        self.assertTrue(math.isclose(
            summary.cell(8, 6).value, rows[-1]['best_val_iou'],
            rel_tol=0.0, abs_tol=1e-14,
        ))
        self.assertIn('不存在 AUC 优先规则', workbook['协议与限制']['A8'].value)
        workbook.close()


if __name__ == '__main__':
    unittest.main()
